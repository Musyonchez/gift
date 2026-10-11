"""Builds the class presentation (PowerPoint) from the outputs of ts_analysis.py. Run: python build_slides.py"""
import json
import numpy as np, pandas as pd
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION

DARK, MOSS, AMBER, INK, MUTE, LIGHT, WHITE, RED = "1E4D2B", "6FA64A", "D98E04", "1F2A22", "5F6B62", "EEF4EA", "FFFFFF", "B03A2E"
HEAD, BODY = "Cambria", "Calibri"
rgb = lambda h: RGBColor.from_string(h)

y = pd.read_csv("series_for_dashboard.csv", parse_dates=[0], index_col=0)["actual"]
fin = pd.read_csv("final_forecast_24m.csv", parse_dates=[0], index_col=0)
hm = pd.read_csv("holdout_metrics.csv").sort_values("RMSE")
rm = pd.read_csv("rolling_metrics.csv").sort_values("mean RMSE")
R = json.load(open("results.json"))
ratio = (y / y.rolling(12, center=True).mean()).dropna()
seas = ratio.groupby(ratio.index.month).mean()

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def bg(slide, color):
    f = slide.background.fill; f.solid(); f.fore_color.rgb = rgb(color)


def text(slide, x, y_, w, h, t, size=16, bold=False, color=INK, font=BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y_), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = anchor
    paras = t if isinstance(t, list) else [t]
    for i, ptxt in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        if i: p.space_before = Pt(8)
        r = p.add_run(); r.text = ptxt; r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic; r.font.name = font; r.font.color.rgb = rgb(color)
    return tb


def title(slide, t, dark=False):
    text(slide, 0.6, 0.45, 12.1, 0.9, t, 34, True, WHITE if dark else DARK, HEAD, anchor=MSO_ANCHOR.MIDDLE)


def card(slide, x, y_, w, h, fill=LIGHT):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y_), Inches(w), Inches(h))
    s.adjustments[0] = 0.06; s.fill.solid(); s.fill.fore_color.rgb = rgb(fill); s.line.fill.background(); s.shadow.inherit = False
    return s


def stat(slide, x, y_, w, h, big, label, fill=LIGHT, bigc=DARK, labc=MUTE):
    card(slide, x, y_, w, h, fill)
    text(slide, x + 0.2, y_ + 0.15, w - 0.4, h * 0.55, big, 34, True, bigc, HEAD, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, x + 0.2, y_ + h * 0.62, w - 0.4, h * 0.34, label, 14, False, labc)


def notes(slide, t): slide.notes_slide.notes_text_frame.text = t


def style_chart(ch, legend=False, size=12):
    ch.font.size = Pt(size); ch.font.name = BODY; ch.font.color.rgb = rgb(MUTE)
    ch.has_legend = legend
    if legend: ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(size)
    try:
        ch.value_axis.major_gridlines.format.line.color.rgb = rgb("DDE5DA"); ch.value_axis.format.line.fill.background()
        ch.category_axis.format.line.color.rgb = rgb("B8C4B6")
    except Exception:
        pass


def skip_labels(ch, n):
    ax = ch.category_axis._element; ns = "http://schemas.openxmlformats.org/drawingml/2006/chart"
    for tag in ("tickLblSkip", "tickMarkSkip"):
        for old in ax.findall(f"{{{ns}}}{tag}"): ax.remove(old)
    nm = ax.find(f"{{{ns}}}noMultiLvlLbl")
    for tag in ("tickLblSkip", "tickMarkSkip"):
        e = etree.SubElement(ax, f"{{{ns}}}{tag}"); e.set("val", str(n))
        if nm is not None: nm.addprevious(e)


# ---------- 1 title
s = prs.slides.add_slide(BLANK); bg(s, DARK)
text(s, 0.8, 2.0, 11.5, 1.8, "Forecasting the Global Tea Price", 52, True, WHITE, HEAD)
text(s, 0.8, 3.7, 11.5, 0.8, "A time series study from a Kenyan export perspective", 24, False, "CFE3C2")
text(s, 0.8, 5.6, 11.5, 0.9, ["STA 3050: Time Series Analysis and Forecasting | Term Paper presentation", "Name: ______________    Student ID: ______________"], 16, False, "CFE3C2")
notes(s, "Introduce the topic: what the tea price is, why it matters to Kenya, and the four questions the study answers. About 30 seconds.")

# ---------- 2 why
s = prs.slides.add_slide(BLANK); bg(s, WHITE); title(s, "Why tea, and why time series")
text(s, 0.6, 1.6, 6.0, 3.6, ["Tea is one of Kenya's largest export earners; Kenya is among the world's leading black tea exporters.",
                             "Farmers, factories and the Treasury all face a world price they cannot control.",
                             "Planning (payments, budgets, contracts) needs a forecast and a measure of its uncertainty."], 18)
for i, (b, l) in enumerate([("1", "How has the price behaved? Is it seasonal?"), ("2", "Is it stationary? What is the autocorrelation?"), ("3", "Which model forecasts best? Does ML win?"), ("4", "What is the 24-month outlook, and how uncertain?")]):
    yy = 1.6 + i * 1.12
    card(s, 7.2, yy, 5.5, 0.95); c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.4), Inches(yy + 0.2), Inches(0.55), Inches(0.55)); c.fill.solid(); c.fill.fore_color.rgb = rgb(DARK); c.line.fill.background()
    c.text_frame.text = b; r = c.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = rgb(WHITE); r.font.name = BODY; c.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    text(s, 8.15, yy + 0.1, 4.4, 0.75, l, 16, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 6.5, 12, 0.5, "Four research questions", 14, False, MUTE, italic=True)
notes(s, "Justify the choice of area (WHY): a commodity the country depends on, monthly data with strong time dependence, and decisions that need forecasts with uncertainty.")

# ---------- 3 data
s = prs.slides.add_slide(BLANK); bg(s, WHITE); title(s, "The data: 34 years of monthly prices")
for i, (b, l) in enumerate([("415", "monthly observations, Jan 1992 to Jul 2026"), ("0", "missing values"), ("256", "mean price, US cents per kg"), ("403", "peak, July 2015")]):
    stat(s, 0.6 + i * 3.1, 1.7, 2.9, 1.9, b, l)
text(s, 0.6, 4.1, 12, 2.6, ["Source: IMF global tea price (average of Mombasa, Colombo and Calcutta auctions), via FRED code PTEAUSDM.",
                            "It is a world benchmark, not the Mombasa price alone, and is nominal in US dollars.",
                            "Low of 143 cents in July 1995; annual average 200 (1992), 346 (2011), 317 (2025); latest value 273 (Jul 2026)."], 17)
notes(s, "Be upfront about what the series is: the IMF benchmark that includes Mombasa, in nominal dollars.")

# ---------- 4 series chart
s = prs.slides.add_slide(BLANK); bg(s, WHITE); title(s, "The price moves in long swings, not around a fixed level")
cd = CategoryChartData(); cd.categories = [d.strftime("%Y-%m") for d in y.index]
cd.add_series("Monthly price", [round(v, 1) for v in y.values]); cd.add_series("12-month average", [None if np.isnan(v) else round(v, 1) for v in y.rolling(12).mean().values])
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(0.6), Inches(1.5), Inches(8.6), Inches(5.3), cd); ch = gf.chart; style_chart(ch, True)
ch.series[0].format.line.color.rgb = rgb(MOSS); ch.series[0].format.line.width = Pt(1.25); ch.series[0].smooth = False
ch.series[1].format.line.color.rgb = rgb(AMBER); ch.series[1].format.line.width = Pt(2.5); ch.series[1].smooth = False
ch.value_axis.has_title = True; ch.value_axis.axis_title.text_frame.text = "US cents per kg"; ch.value_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
skip_labels(ch, 60)
card(s, 9.5, 1.7, 3.3, 3.6)
text(s, 9.75, 1.95, 2.8, 4.4, ["Rise to a 2011 high and a July 2015 peak of 403.", "Strong trend: STL trend strength 0.90.", "Volatility is high: about 23% a year in monthly changes, and it clusters."], 16)
notes(s, "Walk the audience through the chart: no fixed mean, long swings. The 12-month average smooths out the monthly noise.")

# ---------- 5 seasonality
s = prs.slides.add_slide(BLANK); bg(s, WHITE); title(s, "Seasonality is real but small")
cd = CategoryChartData(); cd.categories = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
cd.add_series("Average seasonal ratio", [round((v - 1) * 100, 1) for v in seas.values])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.5), Inches(8.0), Inches(5.3), cd); ch = gf.chart; style_chart(ch)
ch.has_title = True; ch.chart_title.text_frame.text = "Price vs yearly average, % by calendar month"; ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(14); ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
pl = ch.plots[0]; pl.gap_width = 60; pl.has_data_labels = True; pl.data_labels.number_format = '0.0'; pl.data_labels.number_format_is_linked = False; pl.data_labels.font.size = Pt(12); pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
ser = pl.series[0]; ser.invert_if_negative = False
for i, v in enumerate(seas.values):
    pt = ser.points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(MOSS if v >= 1 else AMBER)
for dpt in ser._element.findall("{http://schemas.openxmlformats.org/drawingml/2006/chart}dPt"):
    inv = dpt.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}invertIfNegative")
    if inv is None:
        inv = etree.SubElement(dpt, "{http://schemas.openxmlformats.org/drawingml/2006/chart}invertIfNegative"); dpt.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}idx").addnext(inv)
    inv.set("val", "0")
ch.category_axis.tick_label_position = __import__("pptx.enum.chart", fromlist=["XL_TICK_LABEL_POSITION"]).XL_TICK_LABEL_POSITION.LOW
stat(s, 8.9, 1.7, 3.9, 1.7, "p < 0.001", "ANOVA across months: the pattern is statistically real")
stat(s, 8.9, 3.6, 3.9, 1.7, "0.14 vs 0.90", "Seasonal vs trend strength (STL)", LIGHT)
text(s, 8.9, 5.55, 3.9, 1.2, "Highest in Jan, Feb, Sep and Oct; lowest in May and June, about 4.5% below the yearly average.", 14, False, INK)
notes(s, "Statistically significant but small. A 12-month periodogram peak is absent; the strongest peak is at 6 months, a harmonic of the twice-yearly highs.")

# ---------- 6 stationarity
s = prs.slides.add_slide(BLANK); bg(s, WHITE); title(s, "The price is non-stationary; one difference fixes it")
rows = [("Series", "ADF p-value", "KPSS", "Verdict"), ("Level", "0.260", "rejects (p <= 0.01)", "non-stationary"), ("Log level", "0.289", "rejects (p <= 0.01)", "non-stationary"),
        ("First difference of log", "< 0.001", "not rejected (p >= 0.10)", "stationary"), ("Seasonal difference (12) of log", "0.0001", "not rejected (p >= 0.10)", "stationary")]
tbl = s.shapes.add_table(len(rows), 4, Inches(0.6), Inches(1.6), Inches(12.1), Inches(2.9)).table
for j, w in enumerate([4.2, 2.0, 3.3, 2.6]): tbl.columns[j].width = Inches(w)
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        c = tbl.cell(i, j); c.text = v; p = c.text_frame.paragraphs[0]; p.runs[0].font.size = Pt(16); p.runs[0].font.name = BODY; p.runs[0].font.bold = (i == 0)
        c.fill.solid(); c.fill.fore_color.rgb = rgb(DARK if i == 0 else (LIGHT if i % 2 else WHITE)); p.runs[0].font.color.rgb = rgb(WHITE if i == 0 else INK)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
for k, (b, l) in enumerate([("d = 1", "difference the log price once in ARIMA"), ("0.16", "lag-1 autocorrelation of monthly changes (lags 2 and 3: -0.20, -0.17)"), ("p < 0.001", "ARCH test: volatility clusters")]):
    stat(s, 0.6 + k * 4.1, 4.85, 3.9, 1.85, b, l)
notes(s, "ADF null: unit root. KPSS null: stationary. Both agree, so d=1. Small autocorrelation in the changes motivates ARMA terms on the differences.")

# ---------- 7 holdout accuracy
s = prs.slides.add_slide(BLANK); bg(s, WHITE); title(s, "Hold-out test: gradient boosting wins on 36 months")
hs = hm.iloc[::-1]
cd = CategoryChartData(); cd.categories = [m.replace(" on log", "").replace(" (damped, mult. seasonal)", "") for m in hs["model"]]; cd.add_series("RMSE", [round(v, 1) for v in hs["RMSE"]])
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(1.5), Inches(8.2), Inches(5.4), cd); ch = gf.chart; style_chart(ch, size=12); ch.has_title = False
pl = ch.plots[0]; pl.gap_width = 45; pl.has_data_labels = True; pl.data_labels.font.size = Pt(12); pl.data_labels.number_format = '0.0'; pl.data_labels.number_format_is_linked = False; pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
ser = pl.series[0]
for i, nm in enumerate(hs["model"]):
    pt = ser.points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(AMBER if nm.startswith("Gradient") else ("9AA79C" if nm.startswith("Naive") else MOSS))
ch.value_axis.has_title = True; ch.value_axis.axis_title.text_frame.text = "RMSE, US cents per kg (lower is better)"; ch.value_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
stat(s, 9.1, 1.7, 3.7, 1.8, "32.7", "RMSE of gradient boosting, vs 47.7 for the naive forecast (-31%)")
text(s, 9.1, 3.8, 3.7, 3.0, ["Train: 1992 to Jul 2023 (379 months). Test: the last 36 months.", "ARIMA(2,1,3) is no better than naive (47.8): changes are only weakly predictable.", "One split, one origin: treat the p-values with caution."], 14)
notes(s, "Models: naive, seasonal naive, drift, SES, Holt, Holt-Winters, ARIMA, SARIMA, random forest, gradient boosting, ridge. ML models use lagged returns and month terms with recursive forecasting.")

# ---------- 8 rolling
s = prs.slides.add_slide(BLANK); bg(s, WHITE); title(s, "Across 8 rolling origins the ranking changes")
rs = rm.iloc[::-1]
cd = CategoryChartData(); cd.categories = list(rs["model"]); cd.add_series("Mean RMSE", [round(v, 1) for v in rs["mean RMSE"]])
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(1.5), Inches(7.6), Inches(5.2), cd); ch = gf.chart; style_chart(ch); ch.has_title = False
pl = ch.plots[0]; pl.gap_width = 45; pl.has_data_labels = True; pl.data_labels.font.size = Pt(12); pl.data_labels.number_format = '0.0'; pl.data_labels.number_format_is_linked = False; pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
for i, nm in enumerate(rs["model"]):
    pt = pl.series[0].points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(AMBER if nm.startswith("Holt-W") else ("9AA79C" if nm.startswith("Naive") else MOSS))
ch.value_axis.has_title = True; ch.value_axis.minimum_scale = 0; ch.value_axis.axis_title.text_frame.text = "Mean 12-month RMSE (lower is better)"; ch.value_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
stat(s, 8.5, 1.7, 4.3, 1.8, "27.1 vs 30.0", "Holt-Winters vs naive: about 10% better")
text(s, 8.5, 3.8, 4.3, 3.0, ["Spread across origins is 16 to 19, far larger than the gaps between models.", "Honest conclusion: modest gains over naive; rankings are tentative.", "Choice of origin matters more than choice of model."], 15)
notes(s, "This is the key honesty slide. One hold-out favoured gradient boosting; rolling origins favour Holt-Winters; the differences are small relative to noise.")

# ---------- 9 diagnostics + forecast chart
s = prs.slides.add_slide(BLANK); bg(s, WHITE); title(s, "Forecast: about 274 cents per kg, with wide uncertainty")
hist = y.loc["2022-01-01":]; fidx = list(fin.index); allidx = list(hist.index) + fidx
cats = [d.strftime("%Y-%m") for d in allidx]; n_h = len(hist)
def padded(hist_vals, fc_vals): return [None if v is None else round(float(v), 1) for v in hist_vals] + [None if v is None else round(float(v), 1) for v in fc_vals]
cd = CategoryChartData(); cd.categories = cats
cd.add_series("Actual", padded(hist.values, [None] * len(fidx)))
cd.add_series("ARIMA forecast", padded([None] * (n_h - 1) + [hist.values[-1]], fin.forecast.values))
cd.add_series("80% low", padded([None] * n_h, fin.lo80.values)); cd.add_series("80% high", padded([None] * n_h, fin.hi80.values))
cd.add_series("Holt-Winters", padded([None] * n_h, fin.holt_winters.values))
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(0.6), Inches(1.5), Inches(8.7), Inches(5.3), cd); ch = gf.chart; style_chart(ch, True, 12)
cols = [("1F2A22", 2.25, None), (RED, 2.75, None), ("E3A7A0", 1.5, "dash"), ("E3A7A0", 1.5, "dash"), (AMBER, 1.75, "sysDot")]
from pptx.enum.dml import MSO_LINE
for sr, (c, w, d) in zip(ch.series, cols):
    sr.format.line.color.rgb = rgb(c); sr.format.line.width = Pt(w); sr.smooth = False
    if d: sr.format.line.dash_style = MSO_LINE.DASH if d == "dash" else MSO_LINE.ROUND_DOT
    sr.marker.style = None
    from pptx.enum.chart import XL_MARKER_STYLE
    sr.marker.style = XL_MARKER_STYLE.NONE
ch.value_axis.has_title = True; ch.value_axis.axis_title.text_frame.text = "US cents per kg"; ch.value_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
skip_labels(ch, 12)
stat(s, 9.6, 1.7, 3.2, 1.55, "274", "ARIMA forecast, flat over 24 months")
stat(s, 9.6, 3.45, 3.2, 1.55, "183 to 411", "80% interval at 24 months")
text(s, 9.6, 5.2, 3.2, 1.6, "Holt-Winters and gradient boosting add a mild seasonal rhythm at about 273 to 295. Use them as scenarios.", 14)
notes(s, "Hold-out coverage of the nominal 80% and 95% intervals was 86% and 100%. Residuals show ARCH effects and mild non-normality, so intervals are approximate.")

# ---------- 10 dashboard
s = prs.slides.add_slide(BLANK); bg(s, WHITE); title(s, "Interactive dashboard")
tabs = [("1", "Overview", "zoomable series, range buttons, annual change"), ("2", "Seasonality", "box plots and year-by-month heat-map"), ("3", "Decomposition", "STL trend, seasonal, remainder"),
        ("4", "Stationarity", "ADF and KPSS table, ACF and PACF"), ("5", "Model comparison", "hold-out and rolling errors; toggle models"), ("6", "Forecast", "fan chart with 80% and 95% intervals")]
for i, (n, a, b) in enumerate(tabs):
    x = 0.6 + (i % 3) * 4.1; yy = 1.7 + (i // 3) * 2.35; card(s, x, yy, 3.9, 2.1)
    text(s, x + 0.25, yy + 0.2, 3.4, 0.5, f"{n}. {a}", 20, True, DARK, HEAD); text(s, x + 0.25, yy + 0.85, 3.4, 1.1, b, 16, False, INK)
text(s, 0.6, 6.55, 12, 0.5, "Tea_Price_Dashboard.html: one self-contained file, opens in any browser, works offline.", 15, False, MUTE, italic=True)
notes(s, "Live demo here: open the dashboard, zoom with the range slider, then toggle models in the hold-out chart.")

# ---------- 11 conclusions
s = prs.slides.add_slide(BLANK); bg(s, DARK); title(s, "Conclusions and recommendations", True)
pts = ["The tea price is trending, non-stationary and only weakly predictable; seasonality is small.",
       "ARIMA(2,1,3) captures the autocorrelation, but beats a naive forecast by little.",
       "Holt-Winters and gradient boosting give 10% to 30% lower errors in the tests, with caution about noise.",
       "Plan around the central forecast of about 274 cents per kg and the wide range of 180 to 410, not one number.",
       "Next steps: add weather, exchange rate and supply variables; use GARCH or bootstrap intervals."]
for i, p_ in enumerate(pts):
    text(s, 0.9, 1.7 + i * 1.0, 11.6, 0.9, p_, 20, False, WHITE)
text(s, 0.9, 6.7, 11, 0.5, "Thank you. Questions?", 20, True, "CFE3C2", HEAD)
notes(s, "Close with the practical message and the limitations: univariate, nominal USD benchmark, one hold-out plus eight origins.")

prs.save("Tea_Price_Presentation.pptx")
print("saved", len(prs.slides._sldIdLst), "slides")
