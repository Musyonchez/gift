import nbformat
from docx import Document
from docx.shared import Pt, Inches
nb = nbformat.read('NLP_Assignment_Notebook.ipynb', 4)
cells = {}
cur = None
for c in nb.cells:
    if c.cell_type == 'markdown' and c.source.startswith('## '):
        cur = c.source[3:].strip()
        continue
    if c.cell_type == 'code' and cur:
        out = ''.join(o.get('text', '') for o in c.outputs if o.output_type == 'stream')
        cells[cur.split()[0] if cur.startswith('Q') else cur] = (c.source, out)
doc = Document()
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)


def code(txt, size=8):
    p = doc.add_paragraph()
    r = p.add_run(txt.strip('\n'))
    r.font.name = 'Consolas'
    r.font.size = Pt(size)


def output(txt):
    if not txt.strip():
        return
    doc.add_paragraph('Output:').runs[0].italic = True
    code(txt, 7.5)


def block(key):
    src, out = cells[key]
    code(src)
    output(out)


def ans(t):
    p = doc.add_paragraph()
    r = p.add_run('Answer: ')
    r.bold = True
    p.add_run(t)


def fig(f, w=5.5):
    doc.add_picture(f, width=Inches(w))


def table(rows, header):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Light Grid Accent 1'
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = h
    for r in rows:
        cs = t.add_row().cells
        for i, v in enumerate(r):
            cs[i].text = str(v)


H = doc.add_heading
doc.add_heading('DSA 4020 Assignment 1: Classifying Customer Complaint Tickets', 0)
doc.add_paragraph('Individual assignment. Seed: random_state=42 throughout. The test file is used once, in Q17. Companion files: NLP_Assignment_Notebook.ipynb (runs top to bottom) and predictions.csv.')
H('Setup (imports, seed, data loading)', 1)
code(cells['Setup'][0])

H('Part A: Understand the data', 1)
H('Q1', 2)
block('Q1')
fig('fig_q1_class_dist.png', 4.8)
ans('The training file has 401 tickets (the brief says "about 400"). mpesa_transaction is the largest class at 139 tickets (34.7%) and fraud_scam the smallest at 40 (10.0%), so the data is not balanced: the largest class is 3.5 times the smallest. The average ticket is 36.1 words (median 36, range 10 to 67); length is similar across categories (31.5 to 37.8 words).')
H('Q2', 2)
block('Q2')
ans('Features that make classification harder, with evidence from the data: (1) Several issues in one ticket. T0394 is labelled mpesa_transaction but also contains "Someone is impersonating ur agent and collecting cash from customers", a fraud sentence; T0180 (fraud_scam) is mostly about "Agent withdrew cash from my account but the float was not updated" and "Fuliza amount was deducted", which sound like M-Pesa/loan problems. (2) Typos and SMS shorthand: "transcation", "ur agent", and "teh" appears 41 times in the corpus, so the same word has several spellings. (3) Code-switching: tickets such as "Nimesahau PIN na akaunti yangu imefungwa" mix Kiswahili with English, and the English tools do not know the Kiswahili words. (4) Filler sentences carrying no category signal ("Kindly respond as soon as possible", "This is very frustrating, I have called customer care three times") are added to tickets of every class, and "customer care" is also part of a genuine fraud pattern ("pretending to be customer care").')
H('Q3', 2)
block('Q3')
ans('A model that always predicts mpesa_transaction gets 34.7% accuracy on the training data, which looks like "a third right" for doing nothing. It is misleading because it never identifies any ticket of the other four departments: its macro-F1 is only 0.103 and its recall on fraud_scam is 0. Accuracy rewards the majority class, so with imbalanced classes we also need macro-averaged metrics and per-class recall.')

H('Part B: Preprocessing', 1)
H('Q4', 2)
block('Q4')
ans('The function lowercases, deletes everything that is not a letter (so punctuation and numbers go), tokenises with NLTK and optionally removes stop words, then stems (Porter) or lemmatises (WordNet). Stemming produces non-words ("pleas", "revers", "paybil", "ke" for KES) while lemmatisation keeps readable words.')
H('Q5', 2)
block('Q5')
table([('none', 357), ('stop words only', 304), ('stop words + stemming', 280), ('stop words + lemmatisation', 295)], ['Setting', 'Vocabulary size'])
ans('Removing stop words shrinks the vocabulary most in a single step (-53 types, 357 to 304), because it deletes a whole list of frequent words ("the", "my", "not" ...). Stemming then removes a further 24 types by merging inflected forms (charged/charges/charge), and lemmatisation only 9, because WordNet only merges forms it recognises as English nouns. The reduction is modest in absolute terms because the data is synthetic with a small, repetitive vocabulary; lowercasing and removing punctuation and numbers had already cut the 447 raw whitespace tokens to 357.')
H('Q6', 2)
block('Q6')
ans('Kiswahili/Sheng words found in the data: nimeshindwa (35), pesa (92), mkopo (14), nimesahau (21), akaunti (21), simu (31), nimenunua (18), imefungwa (21), hazijaingia (18), kutuma (13), mtu (13), kuelewa (35). Effect: none of them is in the English stop list, so English stop-word removal leaves them in, including Kiswahili function words that carry no meaning ("na" 56 times, "yangu" 39, "lakini" 40), which then behave like noise. The Porter stemmer and WordNet lemmatiser return every one of them unchanged (for example nimeshindwa -> nimeshindwa, akaunti -> akaunti), so a Kiswahili word is never merged with its inflected forms (nimeshindwa / ameshindwa / tumeshindwa remain three separate features). Two practical solutions: (1) build a Kiswahili/Sheng stop-word list and a small normalisation dictionary (akaunti = account, pesa = money, mkopo = loan, plus typo/shorthand variants such as teh, ur) and apply it before the English pipeline, possibly with language identification per sentence; (2) use character n-gram TF-IDF or fastText-style subword embeddings, which do not depend on a language-specific dictionary and are robust to typos, shorthand and code-switching, trained on the company\'s own multilingual tickets.')
H('Q7', 2)
block('Q7')
ans('Stop-word removal with the NLTK list can hurt here because the list contains negations ("not", "no", "nor", "didn", "wasn", "haven", ...). In complaints the negation is the signal: "payment did NOT reflect", "I do NOT remember making this transaction". Removing them turns "payment did not reflect" into "payment reflect". Stemming also corrupts domain terms ("paybill" -> "paybil", "KES" -> "ke"). In the experiment (5-fold CV, TF-IDF + Logistic Regression) keeping stop words gave macro-F1 0.9094 and removing them 0.9009, a drop of 0.0085. That is the direction I argued, but it is far smaller than the fold-to-fold standard deviation (about 0.03), so on 401 tickets I cannot claim it is a real harm: stop-word removal is at best useless here and risky on real data. Keeping digits (amounts, ids) gave 0.9071, no gain.')

H('Part C: Representation', 1)
H('Q8', 2)
block('Q8')
ans('Both matrices are 401 x 295 (295 lemmatised, stop-word-free tokens) with 8,284 non-zero cells, so they are 93.0% sparse; BoW stores counts and TF-IDF re-weights the same cells. The top terms make sense: account_access: pin, app, account, update; fraud_scam: someone, link, fake, impersonating; loan_credit: loan, repaid, outstanding; mpesa_transaction: transaction, paybill, pesa, kes, payment; network_data: bundle, network, call, hour. A few weaker terms ("keep" for network_data, "customer", "still") come from filler phrases that repeat inside one class ("calls keep dropping"), which shows the data is synthetic and templated.')
H('Q9', 2)
block('Q9')
ans('Unigrams + bigrams scored macro-F1 0.9107 +/- 0.0238 versus 0.8980 +/- 0.0364 for unigrams (Logistic Regression), a gain of 0.013 that is smaller than the standard deviation, so it is suggestive and not conclusive; the matrix grows from 295 to 1,422 features. Example: "paybill" alone points to mpesa_transaction (58 of 72 tickets containing it) and "fake" alone is ambiguous (18 fraud, 11 mpesa), but the bigram "fake paybill" is split almost evenly between fraud_scam (7) and mpesa_transaction (6), so it identifies the scam-paybill pattern that neither unigram expresses; likewise "link text" (fraud 12 of 17) captures the phishing phrase. Honestly, these bigrams are infrequent and the evidence is limited.')
H('Q10', 2)
block('Q10')
ans('Word2Vec (skip-gram, 50 dimensions) was trained on the 401 training tickets only: 357 word types and 14,500 tokens. GloVe was not loaded because I had no pre-trained vectors offline, so only Word2Vec is reported. Some neighbours are sensible (network -> coverage, down, area; pin -> blocked, akaunti, nimesahau, imefungwa; otp -> arrives, codes) but others are noise (loan -> still, what; paybill -> sent, report). (a) A small corpus gives weak embeddings because each word is seen only a few dozen times and the templated sentences repeat, so the model learns co-occurrence in the templates and not general meaning; rare words (e.g. "mkopo", 14 occurrences) are poorly placed. (b) Averaging loses word order and negation (reordered words give an identical vector, shown above), gives filler words the same weight as key words, and squeezes everything into 50 dimensions, so a short fraud sentence inside a long ticket is diluted.')

H('Part D: Modelling', 1)
H('Q11', 2)
block('Q11')
table([('TF-IDF + MultinomialNB', '0.9073', '0.0247'), ('TF-IDF + Logistic Regression', '0.8980', '0.0364'), ('TF-IDF + LinearSVC', '0.9132', '0.0344'), ('TF-IDF (1-2 grams) + LinearSVC', '0.9215', '0.0312'), ('Word2Vec average + Logistic Regression', '0.8934', '0.0500')], ['Model (5-fold stratified CV, training data only)', 'Macro-F1 mean', 'Std'])
ans('All TF-IDF models are within about 0.025 of each other, which is below their standard deviations. The Word2Vec model is worst and most variable (0.8934 +/- 0.0500) and about 30 times slower (the embedding is retrained inside every fold so there is no leakage). Best mean: LinearSVC on 1-2-gram TF-IDF (0.9215).')
H('Q12', 2)
block('Q12')
ans('Single change: class_weight=None versus "balanced" for Logistic Regression on the same folds. With the seed-42 split, balanced scores 0.9007 +/- 0.0316 against 0.8980 +/- 0.0364, a difference of +0.0027 with paired fold differences from -0.016 to +0.023 (std 0.017). Repeating the 5-fold CV with five different shuffles (25 folds) gives 0.8958 versus 0.8963: a mean paired difference of -0.0004 with standard error 0.0048. The difference is far smaller than the noise, so I do not claim either option is better on macro-F1. (I still use "balanced" in the final model as a cheap safeguard for the minority fraud class, but that is a judgement, not something the data proves.)')
H('Q13', 2)
ans('Final model: TF-IDF (unigrams + bigrams) with a Linear SVM (class_weight="balanced"). It has the best mean CV macro-F1 (0.9215 for the unbalanced version), but because the TF-IDF models differ by less than one standard deviation I also choose it on other grounds: it trains in about 0.1 s on 320 tickets, has essentially one hyper-parameter, is nearly as simple as Naive Bayes, and its weights per word/bigram can be read and audited ("link text" pushes towards fraud), which matters when a company must explain routing decisions. The Word2Vec model is about 30x slower, less accurate and opaque. Caveat: I did not tune C or the n-gram range, so I do not claim this is the best possible setting.')

H('Part E: Evaluation and test predictions', 1)
H('Q14', 2)
block('Q13/14')
ans('Trained on 320 tickets and tested on a stratified 81-ticket validation set: accuracy 0.914, macro precision 0.899, macro recall 0.883, macro-F1 0.890, weighted-F1 0.911. loan_credit and mpesa_transaction are almost perfect; account_access (F1 0.741) is the weakest. With only 81 tickets one ticket moves a class recall by several points, so these numbers are noisy; the 5-fold CV over all 401 tickets (output of the Q15 cell) gives macro-F1 0.916 and accuracy 0.925, consistent with this.')
H('Q15', 2)
block('Q15')
fig('fig_q15_confusion.png', 5)
ans('The most frequent confusions on the validation set are account_access predicted as network_data (2) and fraud_scam predicted as account_access (2); in the larger 5-fold CV, account_access is also confused with mpesa_transaction and network_data (4 each). A real misclassified ticket (index 133, true account_access, predicted network_data): "I lost my phone and need to block the line and recover my wallet. Calls fail with a busy tone and the signal bars disappear in the evening." The ticket contains two issues: the second sentence has many network words (calls, signal, busy tone) while the first has few account words, so the model, which only counts words, sees more network evidence. Another (index 109): the customer talks about a blocked PIN (account words) before "Someone called pretending to be customer care and asked for my PIN", so the fraud signal is a small part of the text.')
H('Q16', 2)
ans('Management should watch recall for fraud_scam: a missed fraud ticket (false negative) is routed to the wrong team and the customer keeps losing money, whereas a false alarm only costs an agent a few minutes to re-route. On the 81-ticket validation set fraud_scam precision is 0.857 and recall 0.750 (6 of 8 caught); in 5-fold CV over all 40 fraud tickets precision is 0.875 and recall 0.875 (35 of 40). So roughly 1 in 8 fraud tickets is still missed, too many to rely on without human review.')
H('Q17', 2)
block('Q17')
ans('predictions.csv has 99 rows (the test file contains 99 tickets, "about 100") with exactly the columns ticket_id and category, and all ticket ids match. The model was retrained on all 401 training tickets and the test file was used only here, once. I cannot know the hidden-label performance; I would expect something near the 0.89 to 0.92 macro-F1 seen in validation and CV.')

H('Part F: Reflection and communication', 1)
H('Q18: Note to the Head of Customer Care', 2)
doc.add_paragraph('Subject: Automatic routing of customer complaints: what it can do and what we recommend\n\nWe built a program that reads a complaint and sends it to one of five teams: M-Pesa transactions, network/data, loans/credit, account access, or fraud. On 401 past tickets it routes about 9 in 10 correctly (92.5% in repeated checks, 91% on a held-out set). It is very reliable for loan and M-Pesa complaints and weaker on account-access tickets, where roughly one in four is sent to the wrong team.\n\nWhat it cannot do: it does not understand a complaint, it matches words and word pairs it has seen before. When a customer describes two problems in one message (for example a blocked PIN and then someone pretending to be customer care) it can pick the wrong one. It is weaker with new spellings, Sheng or Kiswahili words that did not appear in the training tickets. Most importantly, about 1 in 8 fraud tickets was missed in our tests, and a missed fraud ticket is the costliest mistake. The tickets we trained on were written examples (synthetic) and only about 400, so real performance may be lower.\n\nRecommendations before deployment: (1) start with the program suggesting a team while a person confirms; (2) automatically send low-confidence tickets (about 1 in 10) and anything mentioning words like PIN, OTP, link or "pretending" to a person, and give fraud tickets a fast lane; (3) retrain every few months with newly labelled tickets, adding Kiswahili/Sheng variants; (4) monitor weekly the share of tickets per team and review a sample of mistakes; (5) strip phone numbers and names from tickets before training.')
H('Q19', 2)
ans('Limitation 1, data: only 401 synthetic, template-like tickets, so the high scores probably overstate real-world accuracy, and the 81-ticket validation set is noisy. With more time and real data I would collect several thousand real tickets, labelled twice with agreement checks, and report confidence intervals. Limitation 2, method/evaluation: hyper-parameters (C, n-gram range, class weight) were not tuned and most model differences were smaller than the CV noise. I would use repeated nested cross-validation, tune on training folds only, and use a multi-label output because tickets mention several issues and some labels are debatable.')
H('Q20', 2)
ans('Risks: (1) privacy leakage: tickets contain phone numbers, names and transaction details that could be exposed from the data set or memorised by a model and its logs; (2) bias and misuse: a model trained on past routing may serve customers who write in Sheng/Kiswahili or with unusual spelling worse, and the data could be reused beyond its purpose (e.g. profiling customers). Safeguards: (1) anonymise before use by masking phone numbers, names, ids and amounts, restrict access by role and store the data encrypted, in line with the Kenya Data Protection Act 2019; (2) keep a human in the loop for low-confidence and fraud tickets, audit accuracy by language group regularly, and set purpose-limitation and retention rules.')
H('Q21 (bonus): Transformer comparison', 2)
ans('Not attempted. My environment has no deep-learning library (torch/transformers) and I did not download pre-trained models, so I cannot report a macro-F1, training time or error pattern for a Transformer, and I will not invent numbers. My expectation, which is untested, is that for 401 template-like tickets with a classical macro-F1 of about 0.92 the extra complexity (GPU, longer training, harder to explain) is probably not justified, though a multilingual sentence-embedding model might help with Kiswahili/Sheng tickets.')

H('Presentation and reproducibility', 1)
doc.add_paragraph('Everything runs from NLP_Assignment_Notebook.ipynb top to bottom with random_state=42; plots have titles and labelled axes; the test file is only used in Q17.')
H('AI-use declaration', 1)
doc.add_paragraph('I used Claude Code (Anthropic, Claude Sonnet 5.5) to write the Python code, run it and draft the written answers in this document. The numbers in the write-up come from the outputs shown. I have reviewed the code and answers and can explain every line. (Edit this declaration to match what you actually did.)')
doc.save('NLP_Assignment_Answers.docx')
